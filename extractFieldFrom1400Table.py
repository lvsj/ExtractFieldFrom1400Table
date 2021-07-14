from docx import Document

NAME_COL = 1
FIELD_NAME_COL = 2
FIELD_TYPE_COL = 3
docFile = 'table.docx'
objectFileName = 'autoGenerate.txt'
txtFile = open(objectFileName, 'w')
# 读入文件
document = Document(docFile)
# 获取文件中的表格集
tables = document.tables
table = tables[0]
for index in range(1, len(table.rows)):
    name = table.cell(index, NAME_COL).text
    fieldName = table.cell(index, FIELD_NAME_COL).text
    fieldType = table.cell(index, FIELD_TYPE_COL).text
    txtFile.write("/** \n")
    txtFile.write(name)
    txtFile.write("\n")
    txtFile.write("*/ \n")
    txtFile.write(" private ")
    if "int" in fieldType:
        txtFile.write(" Integer ")
    else:
        txtFile.write(" String ")
    txtFile.write(fieldName)
    txtFile.write("; \n")
txtFile.close()
